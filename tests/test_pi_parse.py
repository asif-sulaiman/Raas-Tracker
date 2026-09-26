"""PI parsing tests: RAAS-format PDF + DOCX extraction, day-first dates, API filter."""
import io

import pytest

from parse_sales import (
    _dedupe_doubled_text,
    _parse_number,
    extract_client_name,
    extract_pi_date,
    extract_pi_number_from_text,
    parse_pi_stream,
)


RAAS_ITEMS = [
    {"qty": "1,000", "item_no": "1001001", "desc": "SAMPLE DETERGENT\nDetergent",
     "hs": "3402.90.10", "price": "2.50", "total": "2,500.00"},
    {"qty": "4,000", "item_no": "1001002", "desc": "SAMPLE ENZYME T220\nEnzyme",
     "hs": "3507.90.90", "price": "2.50", "total": "10,000.00"},
]


def _make_raas_docx() -> io.BytesIO:
    """Build an in-memory .docx shaped like the RAAS proforma invoice."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("DEMO SUPPLIER SDN. BHD.")
    doc.add_paragraph("Mailing Address")
    doc.add_paragraph("EXAMPLE CLIENT LTD")
    doc.add_paragraph("12, SAMPLE STREET")
    doc.add_paragraph("SAMPLE CITY, SAMPLE COUNTRY")
    doc.add_paragraph("Invoice Number")
    doc.add_paragraph("99000001")
    doc.add_paragraph("Invoice Date")
    doc.add_paragraph("15/09/2026")
    table = doc.add_table(rows=1, cols=7)
    table.rows[0].cells[0].text = "Sr.\nNo."
    table.rows[0].cells[1].text = "Quantity\nin Kg"
    table.rows[0].cells[2].text = "Item No."
    table.rows[0].cells[3].text = "Description"
    table.rows[0].cells[4].text = "HS Code"
    table.rows[0].cells[5].text = "Unit Price\nin USD"
    table.rows[0].cells[6].text = "Total in USD"
    for i, it in enumerate(RAAS_ITEMS, start=1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = it["qty"]
        cells[2].text = it["item_no"]
        cells[3].text = it["desc"]
        cells[4].text = it["hs"]
        cells[5].text = it["price"]
        cells[6].text = it["total"]
    total_row = table.add_row().cells
    total_row[5].text = "Total USD"
    total_row[6].text = "12,500.00"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _make_raas_pdf() -> io.BytesIO:
    """Build an in-memory text-based PDF with a ruled items table (fitz)."""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    y = 40
    for line in ["DEMO SUPPLIER SDN. BHD.", "Mailing Address",
                 "EXAMPLE CLIENT LTD", "12, SAMPLE STREET",
                 "SAMPLE CITY, SAMPLE COUNTRY",
                 "Invoice Number", "99000001",
                 "Invoice Date", "15/09/2026"]:
        page.insert_text((40, y), line, fontsize=10)
        y += 14
    cols = [40, 70, 140, 210, 360, 440, 500, 575]
    rows = [["Sr. No.", "Quantity in Kg", "Item No.", "Description",
             "HS Code", "Unit Price", "Total in USD"]]
    for i, it in enumerate(RAAS_ITEMS, start=1):
        rows.append([str(i), it["qty"], it["item_no"],
                     it["desc"].replace("\n", " "), it["hs"],
                     it["price"], it["total"]])
    rows.append(["", "", "", "", "", "Total USD", "12,500.00"])
    y += 10
    row_h = 22
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            rect = pymupdf.Rect(cols[c], y, cols[c + 1], y + row_h)
            page.draw_rect(rect, color=(0, 0, 0), width=0.5)
            page.insert_textbox(rect + pymupdf.Rect(2, 2, -2, -2), val, fontsize=8)
        y += row_h
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _assert_raas(extraction):
    assert extraction.header.pi_number == "99000001"
    assert extraction.header.pi_date == "2026-09-15"
    assert extraction.header.client_name == "EXAMPLE CLIENT LTD"
    assert extraction.is_valid()
    assert len(extraction.items) == 2
    first, second = extraction.items
    assert first.quantity == 1000.0
    assert first.unit_price == 2.50
    assert first.item_no == "1001001"
    assert first.line_total == 2500.0
    assert "SAMPLE DETERGENT" in first.product_name
    assert second.quantity == 4000.0
    assert second.unit_price == 2.50
    assert second.item_no == "1001002"
    assert second.line_total == 10000.0
    assert "SAMPLE ENZYME" in second.product_name


def test_raas_docx_extraction():
    extraction = parse_pi_stream(_make_raas_docx(), "sample.docx")
    _assert_raas(extraction)
    # Filename carries no PI number, so the body-source notice is expected.
    assert extraction.warnings == ["PI number extracted from document body (not filename)"]


def test_price_column_inferred_when_header_blank():
    from parse_sales import extract_items_from_tables
    tables = [[["Sr.", "Qty", "Item", "Description", "HS", "", "Total"],
               ["1", "1,000", "1001001", "SAMPLE DETERGENT", "3402.90.10",
                "2.50", "2,500.00"]]]
    warnings: list = []
    items = extract_items_from_tables(tables, warnings)
    assert len(items) == 1
    assert items[0].unit_price == 2.50
    assert any("recognized" in w for w in warnings)


def test_raas_pdf_extraction():
    extraction = parse_pi_stream(_make_raas_pdf(), "99000001.pdf")
    _assert_raas(extraction)
    assert extraction.warnings == ["PI number extracted from document body (not filename)"]


def test_dayfirst_slash_dates():
    assert extract_pi_date("Invoice Date\n15/09/2026") == "2026-09-15"
    # Ambiguous 05/09/2026 must read day-first (5 September), not May 9.
    assert extract_pi_date("dated 05/09/2026") == "2026-09-05"
    # Labeled invoice date wins over an earlier stray date.
    assert extract_pi_date("ref 2026-08-15\nInvoice Date\n15/09/2026") == "2026-09-15"


def test_pi_number_from_invoice_label():
    assert extract_pi_number_from_text("Invoice Number\n99000001") == "99000001"
    assert extract_pi_number_from_text("Invoice No. 99000001") == "99000001"
    assert extract_pi_number_from_text("PI Number 99000001") == "99000001"
    assert extract_pi_number_from_text("Proforma Invoice No. 99000001") == "99000001"
    assert extract_pi_number_from_text("PI-2026-001") == "PI-2026-001"


def test_mailing_address_client_and_dedupe():
    assert extract_client_name("Mailing Address\nEXAMPLE CLIENT LTD\n12 SAMPLE") == \
        "EXAMPLE CLIENT LTD"
    assert _dedupe_doubled_text("EXAMPLE CLIENT LTD EXAMPLE CLIENT LTD") == \
        "EXAMPLE CLIENT LTD"
    assert extract_client_name(
        "Mailing Address Delivery Address\nEXAMPLE CLIENT LTD EXAMPLE CLIENT LTD"
    ) == "EXAMPLE CLIENT LTD"


def test_header_row_labels_not_mistaken_for_values():
    # Labels share one header row, values sit on the row below (RAAS-format PDF).
    text = ("Mailing Address Delivery Address Invoice Number Invoice Date\n"
            "EXAMPLE CLIENT LTD EXAMPLE CLIENT LTD 99000001 15/09/2026")
    assert extract_client_name(text) == "EXAMPLE CLIENT LTD"
    assert extract_pi_number_from_text(text) == "99000001"
    assert extract_pi_date(text) == "2026-09-15"


def test_pi_number_value_row_skips_dates():
    # A date on the value row must not be mistaken for the PI number.
    assert extract_pi_number_from_text(
        "Invoice Number\nInvoice Date\n15/09/2026") is None
    assert extract_pi_number_from_text(
        "Invoice Number Invoice Date\n99000001 15/09/2026") == "99000001"


def test_real_client_names_not_label_like():
    from parse_sales import _looks_like_label
    assert not _looks_like_label("GENERIC CHEMICALS LTD")
    assert not _looks_like_label("EXAMPLE CLIENT LTD")
    assert _looks_like_label("Invoice Number")
    assert _looks_like_label("Delivery Address")


def test_delivery_address_fallback():
    assert extract_client_name("Delivery Address\nEXAMPLE CLIENT LTD\n12 SAMPLE") == \
        "EXAMPLE CLIENT LTD"


def test_recover_pair_from_merged_cells():
    from parse_sales import extract_items_from_tables
    tables = [[["Sr.", "Qty", "Item", "Description", "HS", "Price", "Total"],
               ["1", "", "1001001", "GENERIC APC Enzyme 11,000 2.65",
                "3507.90.90", "", "29,150.00"]]]
    warnings: list = []
    items = extract_items_from_tables(tables, warnings)
    assert len(items) == 1
    assert items[0].quantity == 11000.0
    assert items[0].unit_price == 2.65
    assert any("filled in" in w for w in warnings)
    assert not any("does not match" in w for w in warnings)


def test_derive_price_from_total():
    from parse_sales import extract_items_from_tables
    tables = [[["Qty", "Description", "Price", "Total"],
               ["11,000", "GENERIC APC Enzyme", "", "29,150.00"]]]
    warnings: list = []
    items = extract_items_from_tables(tables, warnings)
    assert items[0].unit_price == 2.65
    assert items[0].line_total == 29150.0
    assert any("filled in" in w for w in warnings)


def test_derive_qty_from_total():
    from parse_sales import extract_items_from_tables
    tables = [[["Qty", "Description", "Price", "Total"],
               ["", "GENERIC APC Enzyme", "2.50", "2,500.00"]]]
    warnings: list = []
    items = extract_items_from_tables(tables, warnings)
    assert items[0].quantity == pytest.approx(1000.0)
    assert items[0].line_total == 2500.0
    assert any("filled in" in w for w in warnings)


def test_recovery_ignores_dates_and_item_numbers():
    from parse_sales import extract_items_from_tables
    # No qty/price tokens anywhere: item no. and Sr. must not form a pair.
    tables = [[["Sr.", "Qty", "Item", "Description", "Price", "Total"],
               ["3", "", "1001002", "ENZYME X", "", "10,000.00"]]]
    warnings: list = []
    items = extract_items_from_tables(tables, warnings)
    assert items[0].quantity == 0.0
    assert items[0].unit_price == 0.0
    assert any("does not match" in w for w in warnings)


def test_select_best_tables_prefers_valid():
    from parse_sales import _select_best_tables
    broken = [[["Sr.", "Qty", "Item", "Description", "Price", "Total"],
               ["1", "", "1001001", "GENERIC APC Enzyme", "", "29,150.00"]]]
    good = [[["Sr.", "Qty", "Item", "Description", "Price", "Total"],
             ["1", "11,000", "1001001", "GENERIC APC Enzyme", "2.65", "29,150.00"]]]
    tables, items, warnings = _select_best_tables([broken, good])
    assert tables == good
    assert items[0].quantity == 11000.0
    assert items[0].unit_price == 2.65


def test_pdf_candidates_include_fitz_tables():
    from parse_sales import _pdf_table_candidates
    text, candidates = _pdf_table_candidates(_make_raas_pdf().getvalue())
    assert "99000001" in text
    assert any(candidates[0])
    assert any(candidates[2])


def test_parse_number_space_thousands():
    assert _parse_number("11 000") == 11000.0
    assert _parse_number("15 000") == 15000.0
    assert _parse_number("1 11,000") is None


def test_recover_pair_space_thousands():
    from parse_sales import extract_items_from_tables
    tables = [[["Sr.", "Qty", "Item", "Description", "HS", "Price", "Total"],
               ["1", "", "1001001", "GENERIC APC Enzyme 11 000 2.65",
                "3507.90.90", "", "29,150.00"]]]
    warnings: list = []
    items = extract_items_from_tables(tables, warnings)
    assert len(items) == 1
    assert items[0].quantity == 11000.0
    assert items[0].unit_price == 2.65
    assert items[0].line_total == 29150.0
    assert any("filled in" in w for w in warnings)


def test_layout_block_tables_ignore_grid_and_notes():
    from parse_sales import _layout_block_tables, extract_items_from_tables
    layout = ("Sr. No.  Quantity  Description  Price  Total\n"
              "1  11,000  GENERIC APC Enzyme  2.65  29,150.00\n"
              "Total USD  29,150.00\n"
              "\n"
              "Payment due within 30 days")
    tables = _layout_block_tables(layout)
    assert len(tables) == 1
    warnings: list = []
    items = extract_items_from_tables(tables, warnings)
    assert len(items) == 1
    assert items[0].product_name == "GENERIC APC Enzyme"
    assert items[0].quantity == 11000.0
    assert items[0].unit_price == 2.65


def test_parse_number_formats():
    assert _parse_number("1,000") == 1000.0
    assert _parse_number("US$1,234.50") == 1234.50
    assert _parse_number("2.65") == 2.65
    assert _parse_number("---") is None
    # Merged multi-number cells must not concatenate into a wrong value.
    assert _parse_number("1 11,000") is None
    assert _parse_number("2.65 29,150.00") is None


def test_total_mismatch_warns():
    tables = [[["Qty", "Description", "Unit Price", "Total"],
               ["10", "Widget", "5.00", "999.00"]]]
    from parse_sales import extract_items_from_tables
    warnings: list = []
    items = extract_items_from_tables(tables, warnings)
    assert len(items) == 1
    assert any("does not match" in w for w in warnings)


def test_generic_pi_filename_still_works():
    buf = _make_raas_docx()
    extraction = parse_pi_stream(buf, "PI-2026-001_Client.docx")
    assert extraction.header.pi_number == "PI-2026-001"
    assert len(extraction.items) == 2


def test_api_parse_accepts_pdf_rejects_doc(admin_client):
    pdf = (io.BytesIO(_make_raas_pdf().getvalue()), "99000001.pdf")
    r = admin_client.post("/api/sales/parse", data={"file": pdf},
                          content_type="multipart/form-data")
    assert r.status_code == 200
    body = r.get_json()
    assert body["header"]["pi_number"] == "99000001"
    assert len(body["items"]) == 2

    docx = (io.BytesIO(_make_raas_docx().getvalue()), "sample.docx")
    r = admin_client.post("/api/sales/parse", data={"file": docx},
                          content_type="multipart/form-data")
    assert r.status_code == 200

    r = admin_client.post("/api/sales/parse",
                          data={"file": (io.BytesIO(b"fake"), "legacy.doc")},
                          content_type="multipart/form-data")
    assert r.status_code == 400
